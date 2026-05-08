"""ES 集群巡检系统 - Flask 后端服务

提供文件上传、任务管理、报告查询等 API 接口
"""

import os
import sys
import json
import uuid
import threading
import time
from datetime import datetime
from pathlib import Path
from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from werkzeug.utils import secure_filename
from loguru import logger

# 添加项目根目录和 backend 目录到路径
sys.path.insert(0, str(Path(__file__).parent.parent))  # 项目根目录 /app
sys.path.insert(0, str(Path(__file__).parent))  # backend 目录 /app/backend

from es_insight.diagnostic_extractor import DiagnosticDataExtractor
from es_insight.parse_check_result import CheckResultParser
from generate_html_report import generate_html_report
from db_server import db_manager, TaskModel, ReportModel
from version import VERSION, VERSION_INFO

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    WORD_GENERATION_AVAILABLE = True
except Exception as e:
    WORD_GENERATION_AVAILABLE = False
    logger.warning(f"python-docx 导入失败，Word 生成功能不可用: {str(e)}")


# ============ 配置 ============
class Config:
    # 项目根目录
    BASE_DIR = Path(__file__).parent.parent
    
    # 上传配置
    UPLOAD_FOLDER = BASE_DIR / 'uploads'
    MAX_CONTENT_LENGTH = 2 * 1024 * 1024 * 1024  # 2GB
    
    # 输出目录
    OUTPUT_FOLDER = BASE_DIR / 'out'
    REPORTS_FOLDER = BASE_DIR / 'reports'
    
    # 允许的文件类型
    ALLOWED_EXTENSIONS = {'zip'}
    
    # 配置目录
    CONFIG_FILE = BASE_DIR / 'configs/checkpoints_thresholds_complete.yaml'


# ============ 应用初始化 ============
app = Flask(__name__)
app.config.from_object(Config)
CORS(app)  # 启用跨域支持

# 确保目录存在
app.config['UPLOAD_FOLDER'].mkdir(parents=True, exist_ok=True)
app.config['OUTPUT_FOLDER'].mkdir(parents=True, exist_ok=True)
app.config['REPORTS_FOLDER'].mkdir(parents=True, exist_ok=True)

# 任务存储已迁移到 SQLite 数据库
# tasks = {}  # 已删除，使用数据库存储


# ============ 工具函数 ============

def _set_chinese_font(run, font_name='微软雅黑', font_size=None, bold=None, color=None, italic=None):
    """
    为文本设置中文字体
    
    Args:
        run: docx.text.run.Run 对象
        font_name: 字体名称，默认微软雅黑
        font_size: 字体大小（Pt）
        bold: 是否加粗
        color: 字体颜色（RGBColor）
        italic: 是否斜体
    """
    # 设置西文字体
    run.font.name = font_name
    
    # 设置中文字体（关键：必须通过 XML 设置东亚字体）
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    
    # 设置其他样式
    if font_size is not None:
        run.font.size = font_size
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.italic = italic


def allowed_file(filename):
    """检查文件扩展名是否允许"""
    return '.' in filename and \
           filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS']


def _add_checks_table(doc, checks, lang='zh-CN'):
    """添加检查项表格到文档
    
    Args:
        doc: Word文档对象
        checks: 检查项列表
        lang: 语言代码 (zh-CN, zh-TW, en-US)
    """
    from backend.i18n.translations import translate, translate_check_message
    
    # 创建检查结果表格
    result_table = doc.add_table(rows=1, cols=3)
    result_table.style = 'Table Grid'
    
    # 表头（使用翻译）
    header_row = result_table.rows[0]
    header_cells = header_row.cells
    header_cells[0].text = translate('report.check_details.check_item', lang)
    header_cells[1].text = translate('report.check_details.status', lang)
    header_cells[2].text = translate('report.check_details.details', lang)
    
    # 设置表头样式
    for cell in header_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                _set_chinese_font(run, bold=True, font_size=Pt(9))
    
    # 添加数据行
    for check in checks:
        check_row = result_table.add_row()
        check_cells = check_row.cells
        
        # 智能翻译检查项名称
        check_item_name = check.get('check_item', 'N/A')
        check_item_name = translate_check_message(check_item_name, lang)
        check_cells[0].text = check_item_name
        
        check_cells[1].text = check.get('status', 'N/A')
        
        # 为刚设置的文本设置中文字体
        for cell in [check_cells[0], check_cells[1]]:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_chinese_font(run)
        
        # 根据状态设置颜色
        status = check.get('status', '').lower()
        if 'critical' in status or 'error' in status or '失败' in status:
            for run in check_cells[1].paragraphs[0].runs:
                _set_chinese_font(run, color=RGBColor(255, 0, 0))
        elif 'warning' in status:
            for run in check_cells[1].paragraphs[0].runs:
                _set_chinese_font(run, color=RGBColor(255, 165, 0))
        else:
            for run in check_cells[1].paragraphs[0].runs:
                _set_chinese_font(run, color=RGBColor(0, 128, 0))
        
        # 详情信息（可能有多个）
        details = check.get('details', [])
        if details:
            detail_texts = []
            for detail in details[:3]:  # 最多显示 3 条详情
                if isinstance(detail, dict):
                    detail_msg = detail.get('message', str(detail))
                    # 智能翻译detail消息
                    detail_msg = translate_check_message(detail_msg, lang)
                    detail_texts.append(detail_msg)
                else:
                    detail_str = str(detail)
                    # 智能翻译detail字符串
                    detail_str = translate_check_message(detail_str, lang)
                    detail_texts.append(detail_str)
            check_cells[2].text = '\n'.join(detail_texts) if detail_texts else 'N/A'
            # 为详情文本设置中文字体
            for paragraph in check_cells[2].paragraphs:
                for run in paragraph.runs:
                    _set_chinese_font(run)
        else:
            # 显示值和阈值信息（检查项名称已经在第一列）
            detail_parts = []
            
            value = check.get('value')
            threshold = check.get('threshold')
            if value is not None:
                detail_parts.append(f"值：{value}")
            if threshold:
                detail_parts.append(f"阈值：{threshold}")
            
            detail_text = ', '.join(detail_parts) if detail_parts else check.get('message', 'N/A')
            # 智能翻译detail文本（会翻译"值："和"阈值："）
            detail_text = translate_check_message(detail_text, lang)
            check_cells[2].text = detail_text
            # 为详情文本设置中文字体
            for paragraph in check_cells[2].paragraphs:
                for run in paragraph.runs:
                    _set_chinese_font(run)
        
        # 设置单元格字体大小
        for cell in check_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_chinese_font(run, font_size=Pt(9))
    
    doc.add_paragraph()


def generate_word_report(report_data: dict, output_path: str, lang: str = 'zh-CN'):
    """
    从 JSON 报告数据生成 Word 文档
    
    Args:
        report_data: 包含 basic_info, alerts, check_results 的字典
        output_path: Word 文档输出路径
        lang: 语言代码 (zh-CN, zh-TW, en-US)，默认中文
    """
    from backend.i18n.translations import translate, translate_check_message
    
    doc = Document()
    
    # 设置文档标题（使用翻译）
    title_text = translate('report.title', lang)
    title = doc.add_heading(title_text, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 为标题设置中文字体
    for run in title.runs:
        _set_chinese_font(run, font_size=Pt(18), bold=True)
    
    # 基本信息
    basic_info = report_data.get('basic_info', {})
    alerts = report_data.get('alerts', {})
    check_results = report_data.get('check_results', {})
    
    # 从嵌套结构中提取基本信息
    cluster_info = basic_info.get('cluster_info', {})
    health_summary = basic_info.get('health_summary', {})
    node_summary = basic_info.get('node_summary', {})
    index_summary = basic_info.get('index_summary', {})
    
    # 添加生成时间（使用翻译）
    p = doc.add_paragraph()
    generated_at_label = translate('report.generated_at', lang)
    run = p.add_run(f"{generated_at_label}: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    _set_chinese_font(run, italic=True)
    
    # 健康评估（使用翻译）
    health_assessment_title = translate('report.sections.health_assessment', lang)
    heading1 = doc.add_heading(health_assessment_title, level=1)
    for run in heading1.runs:
        _set_chinese_font(run)
    
    alert_summary = alerts.get('summary', {})
    
    # 优先使用健康评分报告中的评分（新算法），回退到旧算法
    health_score_report = report_data.get('health_score_report', {})
    if health_score_report:
        health_score = health_score_report.get('total_score', alert_summary.get('health_score', 0))
    else:
        health_score = alert_summary.get('health_score', 0)
    
    critical_count = alert_summary.get('critical', 0)
    warning_count = alert_summary.get('warning', 0)
    
    # 健康分数（使用翻译）
    score_para = doc.add_paragraph()
    score_label = translate('report.health.score', lang)
    run1 = score_para.add_run(f"{score_label}: ")
    _set_chinese_font(run1, bold=True)
    score_run = score_para.add_run(f"{health_score}")
    _set_chinese_font(score_run, bold=True, font_size=Pt(14), 
                     color=RGBColor(0, 128, 0) if health_score >= 90 else 
                           (RGBColor(255, 165, 0) if health_score >= 70 else RGBColor(255, 0, 0)))
    
    # 从健康评分详细报告中获取异常问题统计
    health_score_report = report_data.get('health_score_report', {})
    if health_score_report:
        summary_info = health_score_report.get('summary', {})
        critical_issues = summary_info.get('critical_issues_count', 0)
        warning_issues = summary_info.get('warning_issues_count', 0)
        total_issues = summary_info.get('total_issues_count', 0)
        level_1_count = summary_info.get('level_1_count', 0)
        level_2_count = summary_info.get('level_2_count', 0)
        level_3_count = summary_info.get('level_3_count', 0)
        
        # 使用翻译
        level_1_label = translate('report.health.level_1_alerts', lang)
        para2 = doc.add_paragraph()
        run2 = para2.add_run(f"{level_1_label}: {level_1_count}")
        _set_chinese_font(run2)
        
        level_2_label = translate('report.health.level_2_alerts', lang)
        para3 = doc.add_paragraph()
        run3 = para3.add_run(f"{level_2_label}: {level_2_count}")
        _set_chinese_font(run3)
        
        level_3_label = translate('report.health.level_3_alerts', lang)
        para4 = doc.add_paragraph()
        run4 = para4.add_run(f"{level_3_label}: {level_3_count}")
        _set_chinese_font(run4)
    else:
        # 如果没有详细报告，使用旧的告警数据
        level_1_count = 0
        level_2_count = 0
        level_3_count = 0
        para2 = doc.add_paragraph()
        run2 = para2.add_run(f"严重告警数: {critical_count}")
        _set_chinese_font(run2)
        
        para3 = doc.add_paragraph()
        run3 = para3.add_run(f"警告数: {warning_count}")
        _set_chinese_font(run3)
    
    doc.add_paragraph()  # 空行
    
    # ===== 告警信息（放在高级分析组件之前）=====
    _add_alerts_to_word(doc, report_data, lang=lang)
    
    doc.add_paragraph()  # 空行
    
    # ===== 高级分析组件（放在集群基本信息之前）=====
    _add_advanced_analysis_to_word(doc, report_data, lang=lang)
    
    # 集群基本信息（使用翻译）
    cluster_basic_info_title = translate('report.sections.cluster_basic_info', lang)
    heading2 = doc.add_heading(cluster_basic_info_title, level=1)
    for run in heading2.runs:
        _set_chinese_font(run)
    
    # 基本信息表格
    table = doc.add_table(rows=0, cols=2)
    table.style = 'Table Grid'
    
    def add_table_row(label_key, value):
        """添加表格行，label_key为翻译键"""
        row = table.add_row()
        label_text = translate(label_key, lang)
        row.cells[0].text = label_text
        row.cells[1].text = str(value) if value else translate('report.common.na', lang)
        # 设置样式和中文字体
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_chinese_font(run, font_size=Pt(10))
    
    # 第1行：集群名称、集群状态
    add_table_row('report.cluster.cluster_name', cluster_info.get('cluster_name'))
    add_table_row('report.cluster.status', cluster_info.get('status'))
    
    # 第2行：节点数量、数据节点数量
    add_table_row('report.cluster.nodes_count', cluster_info.get('nodes_total'))
    add_table_row('report.cluster.data_nodes', cluster_info.get('data_nodes'))
    
    # 第3行：主节点数量、索引数量
    add_table_row('report.cluster.master_nodes', cluster_info.get('master_nodes'))
    add_table_row('report.cluster.indices_count', cluster_info.get('indices_count'))
    
    # 第4行：分片数量、集群版本
    add_table_row('report.cluster.shards_total', cluster_info.get('shards_total'))
    add_table_row('report.cluster.version', cluster_info.get('es_version'))
    
    # 第5行：许可证书、访问加密认证
    add_table_row('report.cluster.license', cluster_info.get('license_type'))
    
    # 翻译security_enabled字段
    security_enabled = cluster_info.get('security_enabled')
    if isinstance(security_enabled, bool):
        security_text = translate('report.status.yes', lang) if security_enabled else translate('report.status.no', lang)
    elif isinstance(security_enabled, str):
        # 如果是字符串，智能翻译
        security_text = translate_check_message(security_enabled, lang)
    else:
        security_text = str(security_enabled) if security_enabled is not None else translate('report.common.na', lang)
    add_table_row('report.cluster.security', security_text)
    
    # 第6行：冷热分布、冷热分布比例
    has_hot_cold = cluster_info.get('has_hot_cold', False)
    hot_cold_text = translate('report.status.yes', lang) if has_hot_cold else translate('report.status.no', lang)
    add_table_row('report.cluster.hot_cold', hot_cold_text)
    
    # 翻译hot_cold_ratio字段
    hot_cold_ratio = cluster_info.get('hot_cold_ratio', 'N/A')
    if hot_cold_ratio and hot_cold_ratio != 'N/A':
        # 如果是中文，智能翻译
        hot_cold_ratio_text = translate_check_message(str(hot_cold_ratio), lang)
    else:
        hot_cold_ratio_text = translate('report.common.na', lang)
    add_table_row('report.cluster.hot_cold_ratio', hot_cold_ratio_text)
    
    doc.add_paragraph()  # 空行
    
    # 节点基本信息（使用翻译）
    node_basic_info_title = translate('report.sections.node_basic_info', lang)
    heading2_5 = doc.add_heading(node_basic_info_title, level=1)
    for run in heading2_5.runs:
        _set_chinese_font(run)
    
    # 获取节点详细信息
    nodes_detail = basic_info.get('nodes_detail', [])
    
    if nodes_detail:
        # 创建节点基本信息表格（12列）
        node_table = doc.add_table(rows=1, cols=12)
        node_table.style = 'Table Grid'
        
        # 设置表头（使用翻译）
        headers = [
            translate('report.nodes.headers.seq', lang),
            translate('report.nodes.headers.name', lang),
            translate('report.nodes.headers.version', lang),
            translate('report.nodes.headers.cpu', lang),
            translate('report.nodes.headers.memory', lang),
            translate('report.nodes.headers.jvm_heap', lang),
            translate('report.nodes.headers.jvm_heap_usage', lang),
            translate('report.nodes.headers.roles', lang),
            translate('report.nodes.headers.uptime', lang),
            translate('report.nodes.headers.shards', lang),
            translate('report.nodes.headers.disk_space', lang),
            translate('report.nodes.headers.disk_usage', lang),
        ]
        header_cells = node_table.rows[0].cells
        for idx, header_text in enumerate(headers):
            header_cells[idx].text = header_text
            # 设置表头样式
            for paragraph in header_cells[idx].paragraphs:
                for run in paragraph.runs:
                    _set_chinese_font(run, font_size=Pt(9), bold=True)
        
        # 填充数据
        for idx, node in enumerate(nodes_detail, 1):
            row = node_table.add_row()
            cells = row.cells
            
            # 格式化角色显示
            roles = node.get('roles', [])
            if isinstance(roles, list):
                roles_display = ', '.join([str(r) for r in roles if r])
            else:
                roles_display = str(roles)
            if not roles_display:
                roles_display = translate('report.common.unknown', lang)
            
            # 处理可能的空值，提供默认值
            cpu_percent = f"{node.get('cpu_percent', 0):.1f}%"
            heap_used_percent = f"{node.get('heap_used_percent', 0):.1f}%"
            disk_used_percent = f"{node.get('disk_used_percent', 0):.1f}%"
            shard_count = str(node.get('shards_count', 0))
            
            # 辅助函数：处理可能为"未知"的字段，进行翻译
            def translate_unknown(value, lang):
                """如果值是'未知'，则返回翻译后的文本"""
                if value == '未知':
                    return translate('report.common.unknown', lang)
                return value
            
            # 填充12列数据（处理硬编码的"未知"值）
            cells[0].text = str(idx)
            cells[1].text = translate_unknown(node.get('name'), lang) or translate('report.common.unknown', lang)
            cells[2].text = translate_unknown(node.get('version'), lang) or translate('report.common.unknown', lang)
            cells[3].text = cpu_percent
            cells[4].text = translate_unknown(node.get('mem_total'), lang) or translate('report.common.unknown', lang)
            cells[5].text = translate_unknown(node.get('jvm_heap_max'), lang) or translate('report.common.unknown', lang)
            cells[6].text = heap_used_percent
            cells[7].text = roles_display
            cells[8].text = translate_unknown(node.get('uptime'), lang) or translate('report.common.unknown', lang)
            cells[9].text = shard_count
            cells[10].text = translate_unknown(node.get('disk_total'), lang) or translate('report.common.unknown', lang)
            cells[11].text = disk_used_percent
            
            # 设置单元格样式和中文字体
            for cell in cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_chinese_font(run, font_size=Pt(9))
    else:
        p = doc.add_paragraph('无节点数据')
        for run in p.runs:
            _set_chinese_font(run)
    
    doc.add_paragraph()  # 空行
    
    # 分片基本信息（放在第四位）
    shard_summaries = report_data.get('shard_summaries', [])
    shard_distribution = report_data.get('shard_distribution', {})
    large_shards = report_data.get('large_shards', [])
    small_shards_info = report_data.get('small_shards_with_multiple_primaries', {})
    if shard_summaries:
        shard_basic_info_title = translate('report.sections.shard_basic_info', lang)
        heading3_5 = doc.add_heading(shard_basic_info_title, level=1)
        for run in heading3_5.runs:
            _set_chinese_font(run)
        
        # 按照单个分片大小从大到小排序，只取前10个
        shard_summaries_sorted = sorted(shard_summaries, key=lambda x: x.get('avg_shard_size_bytes', 0), reverse=True)[:10]
        
        # 创建分片基本信息表格（4列）
        shard_table = doc.add_table(rows=1, cols=4)
        shard_table.style = 'Table Grid'
        
        # 设置表头（使用翻译）
        headers = [
            translate('report.shards.headers.seq', lang),
            translate('report.shards.headers.index_name', lang),
            translate('report.shards.headers.shard_count', lang),
            translate('report.shards.headers.shard_size', lang),
        ]
        header_cells = shard_table.rows[0].cells
        for idx, header_text in enumerate(headers):
            header_cells[idx].text = header_text
            # 设置表头样式
            for paragraph in header_cells[idx].paragraphs:
                for run in paragraph.runs:
                    _set_chinese_font(run, font_size=Pt(9), bold=True)
        
        # 填充数据
        for idx, summary in enumerate(shard_summaries_sorted, 1):
            row = shard_table.add_row()
            cells = row.cells
            
            # 使用人类可读的格式显示大小（单个分片的平均大小）
            from es_insight.parsers.shards_parser import format_size_human_readable
            size_display = format_size_human_readable(summary.get('avg_shard_size_bytes', 0))
            
            cells[0].text = str(idx)
            cells[1].text = summary.get('index_name', translate('report.common.unknown', lang))
            cells[2].text = str(summary.get('primary_shard_count', 0))
            cells[3].text = size_display
            
            # 设置单元格样式和中文字体
            for cell in cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_chinese_font(run, font_size=Pt(9))
        
        doc.add_paragraph()  # 空行
        
        # 添加主分片分布表格
        if shard_distribution:
            distribution = shard_distribution.get('distribution', {})
            if distribution:
                distribution_title = translate('report.shards.distribution_title', lang)
                doc.add_heading(distribution_title, level=2)
                # 创建主分片分布表格（2列）
                dist_table = doc.add_table(rows=1, cols=2)
                dist_table.style = 'Table Grid'
                
                # 设置表头（使用翻译）
                dist_headers = [
                    translate('report.shards.headers.shard_count', lang),
                    translate('report.shards.headers.index_count', lang)
                ]
                dist_header_cells = dist_table.rows[0].cells
                for idx, header_text in enumerate(dist_headers):
                    dist_header_cells[idx].text = header_text
                    # 设置表头样式
                    for paragraph in dist_header_cells[idx].paragraphs:
                        for run in paragraph.runs:
                            _set_chinese_font(run, font_size=Pt(9), bold=True)
                
                # 填充数据
                for shard_count, index_count in distribution.items():
                    row = dist_table.add_row()
                    cells = row.cells
                    cells[0].text = str(shard_count)
                    cells[1].text = str(index_count)
                    
                    # 设置单元格样式和中文字体
                    for cell in cells:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                _set_chinese_font(run, font_size=Pt(9))
                
                doc.add_paragraph()  # 空行
        
        # 添加单个分片大于40GB索引表格
        large_shards_title = translate('report.shards.large_shards_title', lang)
        doc.add_heading(large_shards_title, level=2)
        # 创建大分片索引表格（3列）
        large_table = doc.add_table(rows=1, cols=3)
        large_table.style = 'Table Grid'
        
        # 设置表头（使用翻译）
        large_headers = [
            translate('report.shards.headers.seq', lang),
            translate('report.shards.headers.index_name', lang),
            translate('report.shards.headers.shard_size', lang)
        ]
        large_header_cells = large_table.rows[0].cells
        for idx, header_text in enumerate(large_headers):
            large_header_cells[idx].text = header_text
            # 设置表头样式
            for paragraph in large_header_cells[idx].paragraphs:
                for run in paragraph.runs:
                    _set_chinese_font(run, font_size=Pt(9), bold=True)
        
        # 填充数据
        if large_shards and len(large_shards) > 0:
            for idx, large_shard in enumerate(large_shards, 1):
                row = large_table.add_row()
                cells = row.cells
                cells[0].text = str(idx)
                cells[1].text = large_shard.get('index_name', '未知')
                cells[2].text = f"{large_shard.get('shard_size_gb', 0):.2f} GB"
                
                # 设置单元格样式和中文字体
                for cell in cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            _set_chinese_font(run, font_size=Pt(9))
        else:
            # 当没有大分片索引时，显示空行
            row = large_table.add_row()
            cells = row.cells
            cells[0].text = ""
            no_data_text = translate('report.shards.no_data', lang)
            cells[1].text = no_data_text
            cells[2].text = ""
            
            # 设置单元格样式和中文字体
            for cell in cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_chinese_font(run, font_size=Pt(9))
        
        doc.add_paragraph()  # 空行
        
        # 添加单个分片小于500MB并且主分片数大于1的索引数量展示
        if small_shards_info:
            count = small_shards_info.get('count', 0)
            small_shards_title = translate('report.shards.small_shards_title', lang)
            doc.add_heading(small_shards_title, level=2)
            p = doc.add_paragraph()
            # 根据语言选择"数量"标签
            if lang == 'en-US':
                count_label = 'Count'
            elif lang == 'zh-TW':
                count_label = '數量'
            else:
                count_label = '数量'
            run = p.add_run(f"{count_label}：{count}")
            _set_chinese_font(run, font_size=Pt(9))
            doc.add_paragraph()  # 空行
    
    # 检查项结果
    check_details_title = translate('report.sections.check_details', lang)
    heading4 = doc.add_heading(check_details_title, level=1)
    for run in heading4.runs:
        _set_chinese_font(run)
    
    if check_results:
        # 跳过 performance 类别（与 HTML 报告保持一致，已删除）
        skipped_categories = {'performance'}
        
        # 为检查类别添加序号
        category_order = [
            'cluster_health',
            'cluster_roles',
            'cluster_destructive',
            'jvm_consistency',
            'cluster_settings',
            'thread_pool_status',
            'performance_stability',
            'primary_shard_rationality',
            'replica_rationality',
            'write_optimization',
            'delayed_allocation',
            'mapping_count',
            'mapping_text_type',
            'segment_overload',
            'swap_partition',
            'index_rejection',
            'search_failure',
            'breaker_status',
            'empty_shards',
            'node_resources',
            'index_health',
            'shards',
            'recovery',
            'ilm',
            'snapshot',
            'security',
            'network',
            'license',
            'node_consistency',
            'index_settings',
            'thread_pools'
        ]
        
        # 使用翻译的检查类别标题（带序号，与HTML报告保持一致）
        category_titles = {
            'cluster_health': f"1、{translate('report.check_details.cluster_health', lang)}",
            'cluster_roles': f"2、{translate('report.check_details.cluster_role', lang)}",
            'cluster_destructive': f"3、{translate('report.check_details.delete_protection', lang)}",
            'jvm_consistency': f"4、{translate('report.check_details.jvm_consistency', lang)}",
            'cluster_settings': f"5、{translate('report.check_details.cluster_params', lang)}",
            'thread_pool_status': f"6、{translate('report.check_details.thread_pool_usage', lang)}",
            'performance_stability': f"7、{translate('report.check_details.gc_status', lang)}",
            'primary_shard_rationality': f"8、{translate('report.check_details.primary_shard_rationality', lang)}",
            'replica_rationality': f"9、{translate('report.check_details.replica_rationality', lang)}",
            'write_optimization': f"10、{translate('report.check_details.write_optimization', lang)}",
            'delayed_allocation': f"11、{translate('report.check_details.delayed_allocation', lang)}",
            'mapping_count': f"12、{translate('report.check_details.mapping_fields', lang)}",
            'mapping_text_type': f"13、{translate('report.check_details.mapping_string_type', lang)}",
            'segment_overload': f"14、{translate('report.check_details.segment_overload', lang)}",
            'swap_partition': f"15、{translate('report.check_details.swap_partition', lang)}",
            'index_rejection': f"16、{translate('report.check_details.index_rejection', lang)}",
            'search_failure': f"17、{translate('report.check_details.search_failure', lang)}",
            'breaker_status': f"18、{translate('report.check_details.parent_circuit_breaker', lang)}",
            'empty_shards': f"19、{translate('report.check_details.empty_shards', lang)}",
            'node_resources': f"20、{translate('report.check_details.node_resources', lang)}",
            'index_health': f"21、{translate('report.check_details.index_health', lang)}",
            'shards': f"22、{translate('report.check_details.shards_check', lang)}",
            'recovery': f"23、{translate('report.check_details.recovery', lang)}",
            'ilm': f"24、{translate('report.check_details.ilm', lang)}",
            'snapshot': f"25、{translate('report.check_details.snapshot', lang)}",
            'security': f"26、{translate('report.check_details.security', lang)}",
            'network': '27、网络配置',  # TODO: 添加翻译键
            'license': '28、许可证状态',  # TODO: 添加翻译键
            'node_consistency': '29、节点一致性',  # TODO: 添加翻译键
            'index_settings': '30、索引设置',  # TODO: 添加翻译键
            'thread_pools': '31、线程池状态'  # TODO: 添加翻译键
        }
        
        # 按照指定顺序处理检查类别
        for category_name in category_order:
            if category_name in check_results:
                category_data = check_results[category_name]
                
                if isinstance(category_data, dict):
                    # 分类标题（带序号）
                    cat_title = category_titles.get(category_name, category_name)
                    cat_heading = doc.add_heading(cat_title, level=2)
                    for run in cat_heading.runs:
                        _set_chinese_font(run)
                    
                    # 检查数据结构：可能是 checks 数组或扁平字典
                    checks = category_data.get('checks', [])
                    
                    if checks:
                        # 情况1：checks 数组结构
                        _add_checks_table(doc, checks, lang=lang)
                    else:
                        # 情况2：扁平字典结构 {"check_key": {"value": ..., "passed": ..., "message": ...}}
                        flat_checks = []
                        for check_key, check_value in category_data.items():
                            if isinstance(check_value, dict):
                                # 使用 detail 字段作为详情，message 作为检查项名称
                                check_item_name = check_value.get('message', check_key)
                                # 智能翻译检查项名称
                                check_item_name = translate_check_message(check_item_name, lang)
                                
                                detail = check_value.get('detail', check_value.get('message', 'N/A'))
                                # 智能翻译detail字段
                                detail = translate_check_message(detail, lang)
                                
                                flat_checks.append({
                                    'check_item': check_item_name,
                                    'status': translate('report.check.status.pass', lang) if check_value.get('passed', True) else translate('report.check.status.fail', lang),
                                    'message': detail,
                                    'value': check_value.get('value'),
                                    'threshold': check_value.get('threshold')
                                })
                        
                        if flat_checks:
                            _add_checks_table(doc, flat_checks, lang=lang)
    
    # 删除原位置的告警信息（已移到健康评分之后）
    # 保存文档
    doc.save(output_path)
    logger.info(f"Word 报告已保存至: {output_path}")


def process_task_async(task_id, zip_path, output_dir):
    """异步处理任务"""
    try:
        logger.info(f"开始处理任务 {task_id}, 文件：{zip_path}")
        
        # 更新任务状态
        TaskModel.update_task_status(task_id, 'processing', progress=10)
        
        # 提取诊断数据（传入 zip_path）
        extractor = DiagnosticDataExtractor(zip_path)
        TaskModel.update_task_status(task_id, 'processing', progress=30)
        
        # 处理 ZIP 文件并保存提取的数据
        result = extractor.extract_all()
        TaskModel.update_task_status(task_id, 'processing', progress=60)
        
        # 直接保存提取的数据到任务目录
        extracted_data_path = Path(output_dir) / f"diagnostic_data_{task_id}.json"
        import json
        with open(extracted_data_path, 'w', encoding='utf-8') as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
        TaskModel.update_task_status(task_id, 'processing', progress=70)
        
        # 解析检查结果（传入提取的数据文件路径和配置文件路径）
        parser = CheckResultParser(
            result_file=str(extracted_data_path),
            config_file=str(app.config['CONFIG_FILE'])
        )
        TaskModel.update_task_status(task_id, 'processing', progress=80)
        
        # 生成报告
        report_id = f"report_{task_id}"
        report_path = app.config['REPORTS_FOLDER'] / f"{report_id}.json"
        
        # 使用 CheckResultParser 生成完整的 JSON 报告（包含 check_results）
        full_report_json = parser.generate_report(output_format='json')
        import json
        full_report_data = json.loads(full_report_json)
        
        # 合并原始提取数据到报告中，供高级分析器使用
        # 添加 nodes_stats 和 nodes 等详细数据
        if 'nodes_stats' in result:
            full_report_data['nodes_stats'] = result['nodes_stats']
        if 'nodes' in result:
            full_report_data['nodes'] = result['nodes']
        if 'nodes_raw' in result:
            full_report_data['nodes_raw'] = result['nodes_raw']  # 原始nodes.json数据，包含settings
        if 'thread_pools' in result:
            full_report_data['thread_pools'] = result['thread_pools']
        if 'shards' in result:
            full_report_data['shards'] = result['shards']
        if 'allocation_explain' in result:
            full_report_data['allocation_explain'] = result['allocation_explain']
        
        # 保存完整报告到 JSON 文件
        with open(report_path, 'w', encoding='utf-8') as f:
            json.dump(full_report_data, f, ensure_ascii=False, indent=2)
        
        # 生成 HTML 报告
        try:
            # full_report_data 已经包含了 generate_html_report 需要的结构：
            # {'basic_info': {...}, 'alerts': {...}, 'check_results': {...}}
            # 异步任务生成时使用默认语言（中文）
            html_content = generate_html_report(full_report_data, lang='zh-CN')
            html_report_path = app.config['REPORTS_FOLDER'] / f"{report_id}.html"
            with open(html_report_path, 'w', encoding='utf-8') as f:
                f.write(html_content)
            logger.info(f"HTML 报告已生成: {html_report_path}")
        except Exception as html_err:
            logger.error(f"生成 HTML 报告失败: {str(html_err)}")
            import traceback
            logger.error(traceback.format_exc())
        
        # 从完整报告中提取前端需要的信息
        basic_info = full_report_data.get('basic_info', {})
        alert_info = full_report_data.get('alerts', {})
        health_score_report = full_report_data.get('health_score_report', {})
        
        # 任务完成时更新数据库
        TaskModel.update_task_status(task_id, 'completed', progress=100)
        
        # 保存报告到数据库
        report_id = f"report_{task_id}"
        ReportModel.create_report(report_id, task_id, full_report_data)
        
        # 使用健康评分报告中的数据（与HTML/Word报告保持一致）
        if health_score_report:
            summary_info = health_score_report.get('summary', {})
            total_issues = summary_info.get('total_issues_count', 0)
            level_1_count = summary_info.get('level_1_count', 0)
            level_2_count = summary_info.get('level_2_count', 0)
            level_3_count = summary_info.get('level_3_count', 0)
        else:
            # 如果没有健康评分报告，回退到旧的数据
            total_issues = alert_info.get('summary', {}).get('critical', 0) + \
                          alert_info.get('summary', {}).get('warning', 0)
            level_1_count = 0
            level_2_count = 0
            level_3_count = 0
        
        tasks_result = {
            'reportId': report_id,
            'healthScore': health_score_report.get('total_score', alert_info.get('summary', {}).get('health_score', 85)),
            'alertCount': total_issues,
            'reportPath': str(report_path),
        }
        
        logger.info(f"任务 {task_id} 处理完成，报告 ID: {report_id}")
        
    except Exception as e:
        logger.error(f"任务 {task_id} 处理失败：{str(e)}")
        TaskModel.update_task_status(task_id, 'failed', error_message=str(e))


# ============ API 路由 ============

@app.route('/api/v1/upload', methods=['POST'])
def upload_file():
    """文件上传接口"""
    try:
        # 检查文件
        if 'file' not in request.files:
            return jsonify({
                'code': 400,
                'message': '未找到上传文件'
            }), 400
        
        file = request.files['file']
        
        if file.filename == '':
            return jsonify({
                'code': 400,
                'message': '未选择文件'
            }), 400
        
        if not allowed_file(file.filename):
            return jsonify({
                'code': 400,
                'message': '只支持 ZIP 格式文件'
            }), 400
        
        # 生成唯一文件名
        original_filename = secure_filename(file.filename)
        unique_filename = f"{uuid.uuid4()}_{original_filename}"
        filepath = app.config['UPLOAD_FOLDER'] / unique_filename
        
        # 保存文件
        file.save(str(filepath))
        
        # 创建任务（保存到数据库）
        task_id = str(uuid.uuid4())
        TaskModel.create_task(task_id, original_filename, os.path.getsize(str(filepath)))
        
        # 启动异步处理线程
        output_dir = app.config['OUTPUT_FOLDER'] / task_id
        output_dir.mkdir(parents=True, exist_ok=True)
        
        thread = threading.Thread(
            target=process_task_async,
            args=(task_id, str(filepath), str(output_dir))
        )
        thread.daemon = True
        thread.start()
        
        logger.info(f"文件上传成功：{original_filename}, 任务 ID: {task_id}")
        
        return jsonify({
            'code': 200,
            'message': '上传成功',
            'data': {
                'taskId': task_id,
                'filename': original_filename,
                'status': 'pending'
            }
        })
        
    except Exception as e:
        logger.error(f"上传失败：{str(e)}")
        return jsonify({
            'code': 500,
            'message': f'上传失败：{str(e)}'
        }), 500


@app.route('/api/v1/tasks', methods=['GET'])
def get_tasks():
    """获取任务列表"""
    try:
        # 从数据库读取任务列表
        task_list = TaskModel.get_all_tasks(limit=100, offset=0)
        
        # 为每个任务添加报告信息
        enriched_tasks = []
        for task in task_list:
            task_dict = dict(task)
            
            # 如果任务已完成，查询关联的报告信息
            if task_dict.get('status') == 'completed':
                with db_manager.get_connection() as conn:
                    cursor = conn.cursor()
                    cursor.execute('''
                        SELECT id, health_score, alert_count, critical_count, warning_count, level_1_count, level_2_count, level_3_count, json_path
                        FROM reports 
                        WHERE task_id = ? AND is_deleted = 0
                        ORDER BY created_at DESC 
                        LIMIT 1
                    ''', (task_dict['id'],))
                    report = cursor.fetchone()
                    
                    if report:
                        # 优先使用数据库中的评分（性能好）
                        health_score = report['health_score']
                        alert_count = report['alert_count']
                        critical_count = report['critical_count'] if report['critical_count'] is not None else 0
                        warning_count = report['warning_count'] if report['warning_count'] is not None else 0
                        level_1_count = report['level_1_count'] if report['level_1_count'] is not None else 0
                        level_2_count = report['level_2_count'] if report['level_2_count'] is not None else 0
                        level_3_count = report['level_3_count'] if report['level_3_count'] is not None else 0
                        
                        task_dict['result'] = {
                            'reportId': report['id'],
                            'healthScore': health_score,
                            'alertCount': alert_count,
                            'criticalCount': critical_count,
                            'warningCount': warning_count,
                            'level_1_count': level_1_count,
                            'level_2_count': level_2_count,
                            'level_3_count': level_3_count
                        }
            
            enriched_tasks.append(task_dict)
        
        return jsonify({
            'code': 200,
            'data': enriched_tasks
        })
        
    except Exception as e:
        logger.error(f"获取任务列表失败：{str(e)}")
        return jsonify({
            'code': 500,
            'message': f'获取任务列表失败：{str(e)}'
        }), 500


@app.route('/api/v1/tasks/<task_id>', methods=['GET'])
def get_task(task_id):
    """获取单个任务详情"""
    try:
        # 从数据库读取任务
        task = TaskModel.get_task(task_id)
        
        if task is None:
            return jsonify({
                'code': 404,
                'message': '任务不存在'
            }), 404
        
        task_dict = dict(task)
        
        # 如果任务已完成，查询关联的报告信息
        if task_dict.get('status') == 'completed':
            with db_manager.get_connection() as conn:
                cursor = conn.cursor()
                cursor.execute('''
                    SELECT id, health_score, alert_count, critical_count, warning_count
                    FROM reports 
                    WHERE task_id = ? AND is_deleted = 0
                    ORDER BY created_at DESC 
                    LIMIT 1
                ''', (task_dict['id'],))
                report = cursor.fetchone()
                
                if report:
                    task_dict['result'] = {
                        'reportId': report['id'],
                        'healthScore': report['health_score'],
                        'alertCount': report['alert_count'],
                        'criticalCount': report['critical_count'] if report['critical_count'] is not None else 0,
                        'warningCount': report['warning_count'] if report['warning_count'] is not None else 0
                    }
        
        return jsonify({
            'code': 200,
            'data': task_dict
        })
        
    except Exception as e:
        logger.error(f"获取任务详情失败：{str(e)}")
        return jsonify({
            'code': 500,
            'message': f'获取任务详情失败：{str(e)}'
        }), 500


@app.route('/api/v1/tasks/<task_id>', methods=['DELETE'])
def delete_task(task_id):
    """删除任务"""
    try:
        # 从数据库获取任务信息
        task = TaskModel.get_task(task_id)
        
        if task is None:
            return jsonify({
                'code': 404,
                'message': '任务不存在'
            }), 404
        
        # 删除相关文件
        # 注意：数据库中不存储 filepath，需要从文件名推断或保留文件
        # 这里我们只删除报告文件
        report_id = f"report_{task_id}"
        
        # 删除报告文件
        for ext in ['.json', '.html', '.docx']:
            report_file = app.config['REPORTS_FOLDER'] / f"{report_id}{ext}"
            try:
                if report_file.exists():
                    os.remove(str(report_file))
            except Exception as e:
                logger.warning(f"删除报告文件失败 {report_file}: {e}")
        
        # 从数据库中删除任务（会级联删除报告记录）
        TaskModel.delete_task(task_id)
        
        logger.info(f"任务已删除：{task_id}")
        
        return jsonify({
            'code': 200,
            'message': '删除成功'
        })
        
    except Exception as e:
        logger.error(f"删除任务失败：{str(e)}")
        return jsonify({
            'code': 500,
            'message': f'删除任务失败：{str(e)}'
        }), 500


@app.route('/api/v1/reports', methods=['GET'])
def get_reports():
    """获取报告列表（支持筛选和分页）"""
    try:
        page = request.args.get('page', 1, type=int)
        per_page = request.args.get('per_page', 20, type=int)
        cluster_name = request.args.get('cluster_name', None)
        min_health = request.args.get('min_health', None, type=float)
        
        offset = (page - 1) * per_page
        reports = ReportModel.get_all_reports(
            limit=per_page, 
            offset=offset,
            cluster_name=cluster_name,
            min_health_score=min_health
        )
        
        return jsonify({
            'code': 200,
            'data': reports,
            'pagination': {
                'page': page,
                'per_page': per_page,
                'total': len(reports)
            }
        })
    except Exception as e:
        logger.error(f"获取报告列表失败：{str(e)}")
        return jsonify({'code': 500, 'message': str(e)}), 500


@app.route('/api/v1/reports/statistics', methods=['GET'])
def get_reports_statistics():
    """获取报告统计信息"""
    try:
        with db_manager.get_connection() as conn:
            cursor = conn.cursor()
            
            # 总报告数
            cursor.execute('SELECT COUNT(*) FROM reports WHERE is_deleted = 0')
            total_reports = cursor.fetchone()[0]
            
            # 平均健康评分
            cursor.execute('SELECT AVG(health_score) FROM reports WHERE is_deleted = 0')
            avg_health = cursor.fetchone()[0] or 0
            
            # 集群列表
            cursor.execute('SELECT DISTINCT cluster_name FROM reports WHERE cluster_name IS NOT NULL AND is_deleted = 0')
            clusters = [row[0] for row in cursor.fetchall()]
        
        return jsonify({
            'code': 200,
            'data': {
                'total_reports': total_reports,
                'avg_health_score': round(avg_health, 2),
                'clusters': clusters
            }
        })
    except Exception as e:
        logger.error(f"获取报告统计信息失败：{str(e)}")
        return jsonify({'code': 500, 'message': str(e)}), 500


@app.route('/api/v1/reports/<report_id>', methods=['GET'])
def get_report(report_id):
    """获取报告详情"""
    try:
        report_path = app.config['REPORTS_FOLDER'] / f"{report_id}.json"
        
        if not report_path.exists():
            return jsonify({
                'code': 404,
                'message': '报告不存在'
            }), 404
        
        import json
        with open(report_path, 'r', encoding='utf-8') as f:
            report_data = json.load(f)
        
        return jsonify({
            'code': 200,
            'data': report_data
        })
        
    except Exception as e:
        logger.error(f"获取报告失败：{str(e)}")
        return jsonify({
            'code': 500,
            'message': f'获取报告失败：{str(e)}'
        }), 500


@app.route('/api/v1/reports/<report_id>/download', methods=['GET'])
def download_report(report_id):
    """下载报告（支持 HTML、JSON、Word）"""
    try:
        # 获取请求参数
        format_type = request.args.get('format', 'html').lower()
        
        # 新增：获取语言参数，默认中文
        lang = request.args.get('lang', 'zh-CN')
        
        # 验证语言是否支持
        from backend.i18n.translations import get_supported_languages
        supported_langs = get_supported_languages()
        if lang not in supported_langs:
            logger.warning(f"不支持的语言: {lang}，使用默认语言 zh-CN")
            lang = 'zh-CN'
        
        logger.info(f"下载报告: {report_id}, 格式: {format_type}, 语言: {lang}")
        
        if format_type == 'word':
            # Word 下载：如果不存在则生成
            # 修改：根据语言生成不同的文件名
            word_filename = f"{report_id}_{lang}.docx" if lang != 'zh-CN' else f"{report_id}.docx"
            word_path = app.config['REPORTS_FOLDER'] / word_filename
            
            if not word_path.exists():
                # 检查是否有 JSON 报告
                json_path = app.config['REPORTS_FOLDER'] / f"{report_id}.json"
                if not json_path.exists():
                    return jsonify({
                        'code': 404,
                        'message': 'JSON 报告不存在，无法生成 Word 文档'
                    }), 404
                
                # 检查 Word 生成是否可用
                if not WORD_GENERATION_AVAILABLE:
                    return jsonify({
                        'code': 503,
                        'message': 'Word 生成功能不可用，请安装 python-docx'
                    }), 503
                
                # 生成 Word 文档（传入语言参数）
                try:
                    logger.info(f"开始生成 Word 报告: {report_id}, 语言: {lang}")
                    import json
                    with open(json_path, 'r', encoding='utf-8') as f:
                        report_data = json.load(f)
                    
                    # 修改：传入语言参数
                    generate_word_report(report_data, str(word_path), lang=lang)
                    logger.info(f"Word 报告已生成: {word_path}")
                except Exception as word_err:
                    logger.error(f"生成 Word 失败: {str(word_err)}")
                    import traceback
                    logger.error(traceback.format_exc())
                    return jsonify({
                        'code': 500,
                        'message': f'生成 Word 失败：{str(word_err)}'
                    }), 500
            
            return send_file(str(word_path), as_attachment=True, download_name=word_filename, mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        
        elif format_type == 'json':
            # JSON 下载
            json_path = app.config['REPORTS_FOLDER'] / f"{report_id}.json"
            if json_path.exists():
                return send_file(str(json_path), as_attachment=True, download_name=f"{report_id}.json")
            else:
                return jsonify({
                    'code': 404,
                    'message': 'JSON 报告不存在'
                }), 404
        
        else:
            # HTML 下载（默认）
            # 检查是否有指定语言的HTML文件
            html_filename = f"{report_id}_{lang}.html" if lang != 'zh-CN' else f"{report_id}.html"
            html_path = app.config['REPORTS_FOLDER'] / html_filename
            
            if not html_path.exists():
                # 如果不存在，尝试从 JSON 重新生成
                json_path = app.config['REPORTS_FOLDER'] / f"{report_id}.json"
                if json_path.exists():
                    try:
                        logger.info(f"开始生成 {lang} 语言的 HTML 报告: {report_id}")
                        import json
                        import sys
                        from pathlib import Path
                        
                        # 确保项目根目录在sys.path中
                        project_root = Path(__file__).parent.parent
                        if str(project_root) not in sys.path:
                            sys.path.insert(0, str(project_root))
                        
                        with open(json_path, 'r', encoding='utf-8') as f:
                            report_data = json.load(f)
                        
                        # 生成指定语言的 HTML 报告
                        from generate_html_report import generate_html_report
                        html_content = generate_html_report(report_data, lang=lang)
                        
                        # 保存生成的 HTML
                        with open(html_path, 'w', encoding='utf-8') as f:
                            f.write(html_content)
                        
                        logger.info(f"{lang} 语言的 HTML 报告已生成: {html_path}")
                    except Exception as html_err:
                        logger.error(f"生成 HTML 失败: {str(html_err)}")
                        import traceback
                        logger.error(traceback.format_exc())
                        # 如果生成失败，尝试返回现有的中文版本
                        html_path = app.config['REPORTS_FOLDER'] / f"{report_id}.html"
                        if not html_path.exists():
                            return jsonify({
                                'code': 500,
                                'message': f'生成 HTML 失败：{str(html_err)}'
                            }), 500
                else:
                    return jsonify({
                        'code': 404,
                        'message': 'JSON 报告不存在，无法生成 HTML'
                    }), 404
            
            return send_file(str(html_path), as_attachment=True, download_name=html_filename)
        
    except Exception as e:
        logger.error(f"下载报告失败：{str(e)}")
        return jsonify({
            'code': 500,
            'message': f'下载报告失败：{str(e)}'
        }), 500


@app.route('/api/v1/reports/<report_id>/view', methods=['GET'])
def view_report(report_id):
    """直接查看 HTML 报告（用于 iframe），支持多语言"""
    try:
        # 获取语言参数，默认为中文
        lang = request.args.get('lang', 'zh-CN')
        
        # 根据语言确定文件路径（与download逻辑保持一致）
        if lang == 'zh-CN':
            html_filename = f"{report_id}.html"
        else:
            html_filename = f"{report_id}_{lang}.html"
        
        html_path = app.config['REPORTS_FOLDER'] / html_filename
        
        # 如果文件不存在，尝试从 JSON 重新生成（类似download逻辑）
        if not html_path.exists():
            json_path = app.config['REPORTS_FOLDER'] / f"{report_id}.json"
            
            if not json_path.exists():
                return jsonify({
                    'code': 404,
                    'message': 'JSON 数据文件不存在'
                }), 404
            
            try:
                logger.info(f"开始生成 {lang} 语言的 HTML 报告: {report_id}")
                import json
                import sys
                from pathlib import Path
                
                # 确保项目根目录在sys.path中
                project_root = Path(__file__).parent.parent
                if str(project_root) not in sys.path:
                    sys.path.insert(0, str(project_root))
                
                with open(json_path, 'r', encoding='utf-8') as f:
                    report_data = json.load(f)
                
                # 生成指定语言的 HTML 报告
                from generate_html_report import generate_html_report
                html_content = generate_html_report(report_data, lang=lang)
                
                # 保存生成的 HTML 文件（后续请求可直接使用）
                with open(html_path, 'w', encoding='utf-8') as f:
                    f.write(html_content)
                
                logger.info(f"HTML 报告已生成并保存: {html_path}")
                
            except Exception as e:
                logger.error(f"生成 HTML 报告失败：{str(e)}")
                return jsonify({
                    'code': 500,
                    'message': f'生成 HTML 报告失败：{str(e)}'
                }), 500
        
        return send_file(str(html_path), mimetype='text/html')
        
    except Exception as e:
        logger.error(f"查看 HTML 报告失败：{str(e)}")
        return jsonify({
            'code': 500,
            'message': f'查看 HTML 报告失败：{str(e)}'
        }), 500


@app.route('/api/v1/reports/<report_id>/html', methods=['GET'])
def get_report_html(report_id):
    """获取报告 HTML 内容"""
    try:
        html_path = app.config['REPORTS_FOLDER'] / f"{report_id}.html"
        
        if not html_path.exists():
            return jsonify({
                'code': 404,
                'message': 'HTML 报告不存在'
            }), 404
        
        with open(html_path, 'r', encoding='utf-8') as f:
            html_content = f.read()
        
        return jsonify({
            'code': 200,
            'data': html_content
        })
        
    except Exception as e:
        logger.error(f"获取 HTML 报告失败：{str(e)}")
        return jsonify({
            'code': 500,
            'message': f'获取 HTML 报告失败：{str(e)}'
        }), 500


@app.route('/api/v1/health', methods=['GET'])
def health_check():
    """健康检查接口"""
    return jsonify({
        'code': 200,
        'message': '服务运行正常',
        'timestamp': datetime.now().isoformat(),
        'version': VERSION,
        'version_info': VERSION_INFO
    })


# ============ 高级分析 API ============

@app.route('/api/v1/reports/<report_id>/thread-pools', methods=['GET'])
def get_thread_pool_analysis(report_id):
    """获取线程池分析结果"""
    try:
        # 加载报告数据
        report_path = app.config['REPORTS_FOLDER'] / f"{report_id}.json"
        with open(report_path, 'r', encoding='utf-8') as f:
            report_data = json.load(f)
        
        from es_insight.analyzers.thread_pool_analyzer import ThreadPoolAnalyzer
        analyzer = ThreadPoolAnalyzer(report_data)
        
        critical_pools = analyzer.get_critical_pools()
        
        return jsonify({
            'code': 200,
            'data': {
                'critical_pools': [p.__dict__ for p in critical_pools],
                'total_pools': len(analyzer.analyze_thread_pools())
            }
        })
    except FileNotFoundError:
        logger.error(f"报告文件不存在: {report_id}")
        return jsonify({'code': 404, 'message': '报告不存在'}), 404
    except Exception as e:
        logger.error(f"获取线程池分析失败: {e}")
        return jsonify({'code': 500, 'message': str(e)}), 500


@app.route('/api/v1/reports/<report_id>/shard-heatmap', methods=['GET'])
def get_shard_heatmap(report_id):
    """获取分片热力图数据"""
    try:
        report_path = app.config['REPORTS_FOLDER'] / f"{report_id}.json"
        with open(report_path, 'r', encoding='utf-8') as f:
            report_data = json.load(f)
        
        from es_insight.analyzers.shard_heatmap_generator import ShardHeatmapGenerator
        generator = ShardHeatmapGenerator(report_data)
        heatmap_data = generator.generate_heatmap_data()
        
        return jsonify({
            'code': 200,
            'data': [h.__dict__ for h in heatmap_data]
        })
    except FileNotFoundError:
        logger.error(f"报告文件不存在: {report_id}")
        return jsonify({'code': 404, 'message': '报告不存在'}), 404
    except Exception as e:
        logger.error(f"获取分片热力图失败: {e}")
        return jsonify({'code': 500, 'message': str(e)}), 500


@app.route('/api/v1/reports/<report_id>/jvm-profile', methods=['GET'])
def get_jvm_memory_profile(report_id):
    """获取 JVM 内存画像"""
    try:
        report_path = app.config['REPORTS_FOLDER'] / f"{report_id}.json"
        with open(report_path, 'r', encoding='utf-8') as f:
            report_data = json.load(f)
        
        from es_insight.analyzers.jvm_memory_analyzer import JVMMemoryAnalyzer
        analyzer = JVMMemoryAnalyzer(report_data)
        profiles = analyzer.analyze_all_nodes()
        
        return jsonify({
            'code': 200,
            'data': [p.__dict__ for p in profiles]
        })
    except FileNotFoundError:
        logger.error(f"报告文件不存在: {report_id}")
        return jsonify({'code': 404, 'message': '报告不存在'}), 404
    except Exception as e:
        logger.error(f"获取 JVM 内存画像失败: {e}")
        return jsonify({'code': 500, 'message': str(e)}), 500


@app.route('/api/v1/reports/<report_id>/config-drift', methods=['GET'])
def get_config_drift(report_id):
    """获取配置漂移检测结果"""
    try:
        report_path = app.config['REPORTS_FOLDER'] / f"{report_id}.json"
        with open(report_path, 'r', encoding='utf-8') as f:
            report_data = json.load(f)
        
        from es_insight.analyzers.config_drift_detector import ConfigurationDriftDetector
        detector = ConfigurationDriftDetector(report_data)
        drifts = detector.detect_drifts()
        
        return jsonify({
            'code': 200,
            'data': [d.__dict__ for d in drifts]
        })
    except FileNotFoundError:
        logger.error(f"报告文件不存在: {report_id}")
        return jsonify({'code': 404, 'message': '报告不存在'}), 404
    except Exception as e:
        logger.error(f"获取配置漂移检测失败: {e}")
        return jsonify({'code': 500, 'message': str(e)}), 500


@app.route('/api/v1/reports/<report_id>/allocation-failures', methods=['GET'])
def get_allocation_failures(report_id):
    """获取分片分配失败分析"""
    try:
        report_path = app.config['REPORTS_FOLDER'] / f"{report_id}.json"
        with open(report_path, 'r', encoding='utf-8') as f:
            report_data = json.load(f)
        
        from es_insight.analyzers.shard_allocation_analyzer import ShardAllocationAnalyzer
        analyzer = ShardAllocationAnalyzer(report_data)
        failures = analyzer.analyze_failures()
        
        return jsonify({
            'code': 200,
            'data': [f.__dict__ for f in failures]
        })
    except FileNotFoundError:
        logger.error(f"报告文件不存在: {report_id}")
        return jsonify({'code': 404, 'message': '报告不存在'}), 404
    except Exception as e:
        logger.error(f"获取分片分配失败分析失败: {e}")
        return jsonify({'code': 500, 'message': str(e)}), 500


# ============ 错误处理 ============
@app.errorhandler(404)
def not_found(error):
    return jsonify({
        'code': 404,
        'message': '接口不存在'
    }), 404


@app.errorhandler(500)
def internal_error(error):
    return jsonify({
        'code': 500,
        'message': '服务器内部错误'
    }), 500


def _add_advanced_analysis_to_word(doc, report_data, lang='zh-CN'):
    """
    向 Word 文档添加高级分析组件（线程池、热力图、JVM、配置漂移、分配状态）
    
    Args:
        doc: python-docx Document 对象
        report_data: 报告数据字典
        lang: 语言代码 (zh-CN, zh-TW, en-US)
    """
    from docx.shared import RGBColor
    from backend.i18n.translations import translate
    
    # 1. 线程池队列堆积告警
    thread_pools = report_data.get('thread_pools', [])
    critical_pools = [p for p in thread_pools if p.get('queue_size', 0) > 1000 or p.get('rejected', 0) > 100]
    
    heading_tp = doc.add_heading(translate('report.advanced.thread_pool_warning', lang), level=2)
    for run in heading_tp.runs:
        _set_chinese_font(run)
    
    if critical_pools:
        table = doc.add_table(rows=len(critical_pools) + 1, cols=5)
        table.style = 'Table Grid'
        
        # 表头（使用翻译）
        headers = [
            translate('report.advanced.thread_pool.headers.node', lang),
            translate('report.advanced.thread_pool.headers.pool', lang),
            translate('report.advanced.thread_pool.headers.queue', lang),
            translate('report.advanced.thread_pool.headers.rejected', lang),
            translate('report.advanced.thread_pool.headers.status', lang)
        ]
        header_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            header_cells[i].text = header
            for paragraph in header_cells[i].paragraphs:
                for run in paragraph.runs:
                    _set_chinese_font(run, font_size=Pt(9), bold=True)
        
        # 数据行
        for idx, pool in enumerate(critical_pools, 1):
            row = table.rows[idx]
            row.cells[0].text = pool.get('node_name', 'N/A')
            row.cells[1].text = pool.get('name', 'N/A')
            row.cells[2].text = str(pool.get('queue_size', 0))
            row.cells[3].text = str(pool.get('rejected', 0))
            
            # 状态（使用翻译）
            queue_size = pool.get('queue_size', 0)
            rejected = pool.get('rejected', 0)
            if rejected > 100:
                row.cells[4].text = translate('report.advanced.thread_pool.overload', lang)
            elif queue_size > 1000:
                row.cells[4].text = translate('report.advanced.thread_pool.high_load', lang)
            else:
                row.cells[4].text = translate('report.advanced.thread_pool.normal', lang)
            
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_chinese_font(run, font_size=Pt(9))
    else:
        if thread_pools:
            p = doc.add_paragraph(translate('report.advanced.thread_pool_normal', lang))
            for run in p.runs:
                _set_chinese_font(run)
        else:
            p = doc.add_paragraph(translate('report.advanced.thread_pool_no_data', lang))
            for run in p.runs:
                _set_chinese_font(run)
    
    doc.add_paragraph()  # 空行
    
    # 2. 分片分布热点
    nodes_stats = report_data.get('nodes_stats', {})
    shards = report_data.get('shards', [])
    nodes = report_data.get('nodes', [])
    
    heading_heatmap = doc.add_heading(translate('report.advanced.shard_heatmap', lang), level=2)
    for run in heading_heatmap.runs:
        _set_chinese_font(run)
    
    if nodes and shards:
        # 计算每个节点的负载指数
        node_shard_count = {}
        for shard in shards:
            node = shard.get('node', '')
            if node:
                node_shard_count[node] = node_shard_count.get(node, 0) + 1
        
        # 只显示前10个节点
        heatmap_data = []
        for node in nodes:
            node_name = node.get('name', '')
            shard_count = node_shard_count.get(node_name, 0)
            cpu = node.get('cpu_percent', 0)
            disk = node.get('disk_used_percent', 0)
            
            # 计算负载指数 (分片40% + CPU 35% + 磁盘25%)
            load_index = min(shard_count / 1000 * 100, 100) * 0.4 + min(cpu, 100) * 0.35 + min(disk, 100) * 0.25
            
            heatmap_data.append({
                'node_name': node_name,
                'load_index': round(load_index, 2),
                'cpu': cpu,
                'disk': disk,
                'shard_count': shard_count
            })
        
        # 按负载指数排序
        heatmap_data.sort(key=lambda x: x['load_index'], reverse=True)
        heatmap_data = heatmap_data[:10]  # 只显示前10个
        
        table = doc.add_table(rows=len(heatmap_data) + 1, cols=5)
        table.style = 'Table Grid'
        
        # 表头（使用翻译）
        headers = [
            translate('report.advanced.shard_heatmap.headers.node', lang),
            translate('report.advanced.shard_heatmap.headers.load', lang),
            translate('report.advanced.shard_heatmap.headers.cpu', lang),
            translate('report.advanced.shard_heatmap.headers.disk', lang),
            translate('report.advanced.shard_heatmap.headers.shards', lang)
        ]
        header_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            header_cells[i].text = header
            for paragraph in header_cells[i].paragraphs:
                for run in paragraph.runs:
                    _set_chinese_font(run, font_size=Pt(9), bold=True)
        
        for idx, data in enumerate(heatmap_data, 1):
            row = table.rows[idx]
            row.cells[0].text = data['node_name']
            row.cells[1].text = f"{data['load_index']:.1f}%"
            row.cells[2].text = f"{data['cpu']:.1f}%"
            row.cells[3].text = f"{data['disk']:.1f}%"
            row.cells[4].text = str(data['shard_count'])
            
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_chinese_font(run, font_size=Pt(9))
    else:
        p = doc.add_paragraph(translate('report.advanced.shard_heatmap.no_data', lang))
        for run in p.runs:
            _set_chinese_font(run)
    
    doc.add_paragraph()  # 空行
    
    # 3. JVM 内存画像
    jvm_profiles = []
    if nodes_stats and isinstance(nodes_stats, dict):
        nodes_dict = nodes_stats.get('nodes', {})
        
        for node_id, stats in nodes_dict.items():
            jvm = stats.get('jvm', {})
            mem = jvm.get('mem', {})
            if mem:
                node_name = stats.get('name', node_id)
                heap_percent = mem.get('heap_used_percent', 0)
                heap_used = mem.get('heap_used_in_bytes', 0)
                heap_max = mem.get('heap_max_in_bytes', 0)
                
                # 格式化字节
                def format_bytes(b):
                    if b == 0: return '0 B'
                    k = 1024
                    sizes = ['B', 'KB', 'MB', 'GB']
                    i = int(__import__('math').floor(__import__('math').log(b) / __import__('math').log(k)))
                    return f"{b / pow(k, i):.2f} {sizes[i]}"
                
                jvm_profiles.append({
                    'node_name': node_name,
                    'heap_percent': heap_percent,
                    'heap_used': format_bytes(heap_used),
                    'heap_max': format_bytes(heap_max)
                })
    
    heading_jvm = doc.add_heading(translate('report.advanced.jvm_profile', lang), level=2)
    for run in heading_jvm.runs:
        _set_chinese_font(run)
    
    if jvm_profiles:
        table = doc.add_table(rows=len(jvm_profiles) + 1, cols=4)
        table.style = 'Table Grid'
        
        # 表头（使用翻译）
        headers = [
            translate('report.advanced.jvm_profile.headers.node', lang),
            translate('report.advanced.jvm_profile.headers.heap_percent', lang),
            translate('report.advanced.jvm_profile.headers.heap_used', lang),
            translate('report.advanced.jvm_profile.headers.heap_max', lang)
        ]
        header_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            header_cells[i].text = header
            for paragraph in header_cells[i].paragraphs:
                for run in paragraph.runs:
                    _set_chinese_font(run, font_size=Pt(9), bold=True)
        
        for idx, profile in enumerate(jvm_profiles, 1):
            row = table.rows[idx]
            row.cells[0].text = profile['node_name']
            row.cells[1].text = f"{profile['heap_percent']:.1f}%"
            row.cells[2].text = profile['heap_used']
            row.cells[3].text = profile['heap_max']
            
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_chinese_font(run, font_size=Pt(9))
    else:
        p = doc.add_paragraph(translate('report.advanced.jvm_profile.no_data', lang))
        for run in p.runs:
            _set_chinese_font(run)
    
    doc.add_paragraph()  # 空行
    
    # 4. 配置漂移检测
    nodes_raw = report_data.get('nodes_raw', {})
    nodes_with_settings = []
    
    if nodes_raw and isinstance(nodes_raw, dict) and 'nodes' in nodes_raw:
        raw_nodes = nodes_raw.get('nodes', {})
        for node_id, node_data in raw_nodes.items():
            settings = node_data.get('settings', {})
            if settings:
                nodes_with_settings.append({
                    'name': node_data.get('name', node_id),
                    'settings': settings
                })
    
    heading_drift = doc.add_heading(translate('report.advanced.config_drift', lang), level=2)
    for run in heading_drift.runs:
        _set_chinese_font(run)
    
    if nodes_with_settings:
        # 简单的配置一致性检查：比较关键配置项
        config_keys = ['path.data', 'path.logs', 'network.host', 'http.port']
        drifts_found = []
        
        for key in config_keys:
            values = {}
            for node in nodes_with_settings:
                value = node['settings'].get(key, translate('report.advanced.config_drift.not_set', lang))
                if value not in values:
                    values[value] = []
                values[value].append(node['name'])
            
            if len(values) > 1:
                drifts_found.append({
                    'key': key,
                    'variations': values
                })
        
        if drifts_found:
            for drift in drifts_found:
                config_key_label = translate('report.advanced.config_drift.config_key', lang)
                p = doc.add_paragraph(f'{config_key_label}: {drift["key"]}')
                for run in p.runs:
                    _set_chinese_font(run, bold=True)
                
                for value, nodes_list in drift['variations'].items():
                    p = doc.add_paragraph(f'  - {value}: {", ".join(nodes_list)}')
                    for run in p.runs:
                        _set_chinese_font(run)
        else:
            p = doc.add_paragraph(translate('report.advanced.config_drift.all_consistent', lang))
            for run in p.runs:
                _set_chinese_font(run)
    else:
        p = doc.add_paragraph(translate('report.advanced.config_drift.no_data', lang))
        for run in p.runs:
            _set_chinese_font(run)
    
    doc.add_paragraph()  # 空行
    
    # 5. 分片分配状态
    allocation_explain = report_data.get('allocation_explain', {})
    
    heading_alloc = doc.add_heading(translate('report.advanced.shard_allocation', lang), level=2)
    for run in heading_alloc.runs:
        _set_chinese_font(run)
    
    has_failures = allocation_explain and isinstance(allocation_explain, dict) and 'index' in allocation_explain
    is_error_response = allocation_explain and isinstance(allocation_explain, dict) and 'error' in allocation_explain
    
    if has_failures:
        index_name = allocation_explain.get('index', 'N/A')
        shard_id = allocation_explain.get('shard', 'N/A')
        
        unassigned_found_label = translate('report.advanced.shard_allocation.unassigned_found', lang)
        shard_label = translate('report.advanced.shard_allocation.shard_label', lang)
        p = doc.add_paragraph(f'{unassigned_found_label}{index_name}{shard_label}{shard_id}')
        for run in p.runs:
            _set_chinese_font(run)
        
        reasons = []
        unassigned_info = allocation_explain.get('unassigned_info', {})
        if unassigned_info:
            unassigned_reason_label = translate('report.advanced.shard_allocation.unassigned_reason', lang)
            unknown_label = translate('report.advanced.shard_allocation.unknown', lang)
            reasons.append(f"{unassigned_reason_label}: {unassigned_info.get('reason', unknown_label)}")
        
        node_decisions = allocation_explain.get('node_allocation_decisions', [])
        for decision in node_decisions:
            node_name = decision.get('node_name', '')
            deciders = decision.get('deciders', [])
            for decider in deciders:
                if decider.get('decision') == 'NO':
                    reasons.append(f"节点 {node_name}: {decider.get('explanation', '')}")
        
        if reasons:
            failure_reasons_label = translate('report.advanced.shard_allocation.failure_reasons', lang)
            p = doc.add_paragraph(failure_reasons_label)
            for run in p.runs:
                _set_chinese_font(run, bold=True)
            
            for reason in reasons:
                p = doc.add_paragraph(f'  • {reason}', style='List Bullet')
                for run in p.runs:
                    _set_chinese_font(run)
    elif is_error_response:
        p = doc.add_paragraph(translate('report.advanced.shard_allocation.all_normal', lang))
        for run in p.runs:
            _set_chinese_font(run)
    else:
        p = doc.add_paragraph(translate('report.advanced.shard_allocation.all_normal_or_no_data', lang))
        for run in p.runs:
            _set_chinese_font(run)
    
    doc.add_paragraph()  # 空行


def _add_alerts_to_word(doc, report_data, lang='zh-CN'):
    """
    向 Word 文档添加告警信息（放在健康评分之后、高级分析组件之前）
    
    Args:
        doc: python-docx Document 对象
        report_data: 报告数据字典
        lang: 语言代码 (zh-CN, zh-TW, en-US)
    """
    from backend.i18n.translations import translate, translate_check_message
    
    alerts = report_data.get('alerts', {})
    health_score_report = report_data.get('health_score_report', {})
    
    # 告警信息
    alert_details = alerts.get('details', [])
    if health_score_report and health_score_report.get('issues', []):
        heading5 = doc.add_heading(translate('report.alerts.title', lang), level=1)
        for run in heading5.runs:
            _set_chinese_font(run)
        
        # 显示健康评分的问题清单
        issues = health_score_report.get('issues', [])
        if issues:
            # 按严重程度分组问题
            # severity_level: 1=不严重(优化建议), 2=一般(建议处理), 3=非常严重(必须处理)
            level_3_issues_list = [issue for issue in issues if issue.get('severity_level') == 3]  # 最严重
            level_2_issues_list = [issue for issue in issues if issue.get('severity_level') == 2]
            level_1_issues_list = [issue for issue in issues if issue.get('severity_level') == 1]  # 不严重
            
            # 列出三级告警（最严重）
            if level_3_issues_list:
                critical_heading = doc.add_heading(translate('report.alerts.level_3', lang), level=2)
                for run in critical_heading.runs:
                    _set_chinese_font(run, color=RGBColor(220, 53, 69))  # 红色
                
                for i, issue in enumerate(level_3_issues_list, 1):
                    issue_para = doc.add_paragraph(style='List Bullet')
                    # 智能翻译description
                    description = translate_check_message(issue.get('description', '未知问题'), lang)
                    issue_run = issue_para.add_run(description)
                    _set_chinese_font(issue_run, color=RGBColor(220, 53, 69))  # 红色
                    
                    # 添加问题值（如果有的话）
                    if issue.get('value') is not None:
                        current_value_label = translate('report.alerts.current_value_label', lang)
                        # 智能翻译value
                        value_text = translate_check_message(str(issue.get('value')), lang)
                        value_run = issue_para.add_run(f" ({current_value_label}: {value_text})")
                        _set_chinese_font(value_run, italic=True, color=RGBColor(128, 128, 128))
                    
                    # 添加建议
                    if issue.get('suggestion'):
                        suggestion_para = doc.add_paragraph(style='Intense Quote')
                        suggestion_label = translate('report.alerts.suggestion_label', lang)
                        # 智能翻译suggestion
                        suggestion_text = translate_check_message(issue.get('suggestion'), lang)
                        suggestion_run = suggestion_para.add_run(f"{suggestion_label}: {suggestion_text}")
                        _set_chinese_font(suggestion_run, italic=True)
            
            # 列出二级告警
            if level_2_issues_list:
                warning_heading = doc.add_heading(translate('report.alerts.level_2', lang), level=2)
                for run in warning_heading.runs:
                    _set_chinese_font(run, color=RGBColor(255, 193, 7))  # 黄色
                
                for i, issue in enumerate(level_2_issues_list, 1):
                    issue_para = doc.add_paragraph(style='List Bullet')
                    # 智能翻译description
                    description = translate_check_message(issue.get('description', '未知问题'), lang)
                    issue_run = issue_para.add_run(description)
                    _set_chinese_font(issue_run, color=RGBColor(255, 193, 7))  # 黄色
                    
                    # 添加问题值（如果有的话）
                    if issue.get('value') is not None:
                        current_value_label = translate('report.alerts.current_value_label', lang)
                        # 智能翻译value
                        value_text = translate_check_message(str(issue.get('value')), lang)
                        value_run = issue_para.add_run(f" ({current_value_label}: {value_text})")
                        _set_chinese_font(value_run, italic=True, color=RGBColor(128, 128, 128))
                    
                    # 添加建议
                    if issue.get('suggestion'):
                        suggestion_para = doc.add_paragraph(style='Intense Quote')
                        suggestion_label = translate('report.alerts.suggestion_label', lang)
                        # 智能翻译suggestion
                        suggestion_text = translate_check_message(issue.get('suggestion'), lang)
                        suggestion_run = suggestion_para.add_run(f"{suggestion_label}: {suggestion_text}")
                        _set_chinese_font(suggestion_run, italic=True)
            
            # 列出一级告警（不严重）
            if level_1_issues_list:
                info_heading = doc.add_heading(translate('report.alerts.level_1', lang), level=2)
                for run in info_heading.runs:
                    _set_chinese_font(run, color=RGBColor(23, 162, 184))  # 蓝色
                
                for i, issue in enumerate(level_1_issues_list, 1):
                    issue_para = doc.add_paragraph(style='List Bullet')
                    # 智能翻译description
                    description = translate_check_message(issue.get('description', '未知问题'), lang)
                    issue_run = issue_para.add_run(description)
                    _set_chinese_font(issue_run, color=RGBColor(23, 162, 184))  # 蓝色
                    
                    # 添加问题值（如果有的话）
                    if issue.get('value') is not None:
                        current_value_label = translate('report.alerts.current_value_label', lang)
                        # 智能翻译value
                        value_text = translate_check_message(str(issue.get('value')), lang)
                        value_run = issue_para.add_run(f" ({current_value_label}: {value_text})")
                        _set_chinese_font(value_run, italic=True, color=RGBColor(128, 128, 128))
                    
                    # 添加建议
                    if issue.get('suggestion'):
                        suggestion_para = doc.add_paragraph(style='Intense Quote')
                        suggestion_label = translate('report.alerts.suggestion_label', lang)
                        # 智能翻译suggestion
                        suggestion_text = translate_check_message(issue.get('suggestion'), lang)
                        suggestion_run = suggestion_para.add_run(f"{suggestion_label}: {suggestion_text}")
                        _set_chinese_font(suggestion_run, italic=True)
    elif alert_details:
        # 只有在没有健康评分问题时，才显示规则引擎告警
        heading5 = doc.add_heading(translate('report.alerts.title', lang), level=1)
        for run in heading5.runs:
            _set_chinese_font(run)
        
        for alert in alert_details:
            alert_para = doc.add_paragraph()
            severity = alert.get('severity', 'INFO')
            
            # 告警级别
            severity_run = alert_para.add_run(f"[{severity}] ")
            _set_chinese_font(severity_run, bold=True, 
                            color=RGBColor(255, 0, 0) if severity == 'CRITICAL' else 
                                  (RGBColor(255, 165, 0) if severity == 'WARNING' else RGBColor(0, 0, 255)))
            
            # 告警消息
            msg_run = alert_para.add_run(alert.get('message', ''))
            _set_chinese_font(msg_run)
            
            # 建议
            if alert.get('suggestion'):
                suggestion_para = doc.add_paragraph()
                sug_run1 = suggestion_para.add_run("建议: ")
                _set_chinese_font(sug_run1, italic=True)
                sug_run2 = suggestion_para.add_run(alert.get('suggestion'))
                _set_chinese_font(sug_run2)
    else:
        # 两者都没有，显示无告警
        heading5 = doc.add_heading('告警信息', level=1)
        for run in heading5.runs:
            _set_chinese_font(run)
        
        no_alert_para = doc.add_paragraph()
        no_alert_run = no_alert_para.add_run("✅ 无告警，系统运行正常")
        _set_chinese_font(no_alert_run)


# ============ 主程序 ============
if __name__ == '__main__':
    # 配置日志
    logger.remove()
    logger.add(sys.stderr, level="INFO")
    
    logger.info("Starting ES Insight Backend Service...")
    logger.info(f"Upload folder: {app.config['UPLOAD_FOLDER']}")
    logger.info(f"Output folder: {app.config['OUTPUT_FOLDER']}")
    logger.info(f"Reports folder: {app.config['REPORTS_FOLDER']}")
    
    # 启动服务
    app.run(
        host='0.0.0.0',
        port=5000,
        debug=False,
        threaded=True
    )
